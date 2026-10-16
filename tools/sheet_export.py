#!/usr/bin/env python3
"""
Build the editable directory workbook from resources.json.

    python3 tools/sheet_export.py resources.json sheet/krisis-directory.xlsx

Upload the .xlsx to Google Drive and open it with Google Sheets (dropdowns
carry over). When you are ready to bring changes back, download it again as
.xlsx and run tools/sheet_import.py.
"""
import json
import sys

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from sheet_common import (ADD_NEW_COLS, CONTACT_TYPES, DIRECTORY_COLS,
                          NATIONAL, PUBLISH, STATUSES, notes_list_to_text)

HEAD = PatternFill("solid", fgColor="1F3A5F")
LOCKED = PatternFill("solid", fgColor="E8E8E8")
MAX_ROWS = 2000

HELP = {
    "Directory": {
        "ID": "Set by the importer. Never type or change an ID.",
        "Category": "Pick from the list.",
        "Council": f"One council, or {NATIONAL} for a national service.",
        "Hours by location": "One per line: Label = text",
        "Publish": "yes = shown on the site; review = hidden pending a check; no = hidden.",
        "Verified": "yes only if every contact detail was checked on the organisation's own site; otherwise no.",
        "Hide contact details": "yes hides phone/email on the site and shows the website link only.",
    },
    "Add new": {
        "Councils served": f"Every council it serves, separated by ; (e.g. Fife; West Lothian). Or {NATIONAL}. Do NOT add one row per council.",
        "Phone type": "helpline / switchboard / referral line / none. A switchboard is never published as a support line.",
        "Age range (their words)": "Copy the age range exactly as their own website states it.",
        "Source page": "The exact page on the organisation's OWN site where you found the details. Directories do not count.",
        "Date checked": "YYYY-MM-DD",
        "Check result": "ok = details confirmed; link-only = real but no public phone/email; dead / moved / unclear are never added.",
        "Phone": "Only a number you saw on their own site. Blank is better than a guess.",
    },
}


def header(ws, cols, tab):
    for i, (label, _, width) in enumerate(cols, 1):
        c = ws.cell(row=1, column=i, value=label)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[get_column_letter(i)].width = width
        tip = HELP.get(tab, {}).get(label)
        if tip:
            c.comment = Comment(tip, "KRISIS")
    ws.freeze_panes = "B2" if tab == "Directory" else "D2"


def dropdown(ws, cols, label, ref):
    idx = [c[0] for c in cols].index(label) + 1
    col = get_column_letter(idx)
    dv = DataValidation(type="list", formula1=ref, allow_blank=True,
                        showErrorMessage=True, errorTitle="Not in list",
                        error="Pick a value from the dropdown.")
    ws.add_data_validation(dv)
    dv.add(f"{col}2:{col}{MAX_ROWS}")


def main(src, dst):
    data = json.load(open(src, encoding="utf-8"))
    cat_label = {c["slug"]: c["label"] for c in data["categories"]}
    councils = [c["name"] for c in data["councils"]]

    wb = Workbook()

    # ---- Read me
    rm = wb.active
    rm.title = "Read me"
    lines = [
        "KRISIS directory: shared working sheet",
        "",
        "Directory tab: every record currently in resources.json, one row each. Correct mistakes here.",
        "Add new tab: ONE ROW PER SERVICE. Put every council it serves in 'Councils served', separated by ;",
        f"  or write {NATIONAL}. Never copy the same service into a row per council.",
        "",
        "Rules (same as the verification protocol):",
        "  - Only the organisation's own website counts as a source. Directories, ALISS, council lists are leads only.",
        "  - Blank is a safe answer. A wrong number is worse than no number.",
        "  - Record the exact page in 'Source page' and the date in 'Date checked'.",
        "  - Rows with Check result dead / moved / unclear, or blank, are never added.",
        "",
        "Do not rename, reorder or delete columns or tabs. Do not edit the ID column.",
        "Hover a column header for help on that column.",
        "",
        f"Built from resources.json generated {data.get('generated', '')} ({len(data['resources'])} records).",
    ]
    for i, t in enumerate(lines, 1):
        rm.cell(row=i, column=1, value=t)
    rm["A1"].font = Font(bold=True, size=14)
    rm.column_dimensions["A"].width = 120

    # ---- Lists (dropdown sources)
    ls = wb.create_sheet("Lists")
    cols = {
        "Categories": [cat_label[c["slug"]] for c in data["categories"]],
        "Councils": [NATIONAL] + councils,
        "Publish": PUBLISH,
        "Yes": ["yes"],
        "Yes/No": ["yes", "no"],
        "Phone type": CONTACT_TYPES,
        "Check result": STATUSES,
    }
    refs = {}
    for ci, (name, values) in enumerate(cols.items(), 1):
        L = get_column_letter(ci)
        ls.cell(row=1, column=ci, value=name).font = Font(bold=True)
        for ri, v in enumerate(values, 2):
            ls.cell(row=ri, column=ci, value=v)
        refs[name] = f"Lists!${L}$2:${L}${len(values) + 1}"
        ls.column_dimensions[L].width = 28

    # importer refuses a sheet built from an older resources.json
    ls["H1"] = "Built from generated"
    ls["H1"].font = Font(bold=True)
    ls["H2"] = data.get("generated", "")
    ls.column_dimensions["H"].width = 24

    # ---- Directory
    d = wb.create_sheet("Directory", 1)
    header(d, DIRECTORY_COLS, "Directory")
    for ri, r in enumerate(data["resources"], 2):
        for ci, (_, key, _) in enumerate(DIRECTORY_COLS, 1):
            if key == "category":
                v = cat_label.get(r.get("category"), r.get("category"))
            elif key == "local_authority":
                v = NATIONAL if r.get("scope") == "national" else r.get("local_authority", "")
            elif key == "notes_list":
                v = notes_list_to_text(r.get("notes_list"))
            else:
                v = r.get(key, "")
            c = d.cell(row=ri, column=ci, value=v if v is not None else "")
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if key == "id":
                c.fill = LOCKED
    for label, ref in [("Category", "Categories"), ("Council", "Councils"),
                       ("Publish", "Publish"), ("Verified", "Yes/No"),
                       ("Hide contact details", "Yes")]:
        dropdown(d, DIRECTORY_COLS, label, refs[ref])
    d.auto_filter.ref = f"A1:{get_column_letter(len(DIRECTORY_COLS))}{len(data['resources']) + 1}"

    # ---- Add new
    a = wb.create_sheet("Add new", 2)
    header(a, ADD_NEW_COLS, "Add new")
    for label, ref in [("Category", "Categories"), ("Phone type", "Phone type"),
                       ("Check result", "Check result")]:
        dropdown(a, ADD_NEW_COLS, label, refs[ref])

    wb.save(dst)
    print(f"wrote {dst}: {len(data['resources'])} directory rows, "
          f"{len(data['categories'])} categories, {len(councils)} councils")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(*sys.argv[1:])
